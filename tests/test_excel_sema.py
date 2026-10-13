"""Arşiv okumasının sütun numarasından koparılması (bilet #4).

Okuma üç kademelidir: gömülü şema sayfası → başlık eşlemesi → sabit sütun
numaraları. Her kademe ayrı test edilir; hepsi gerçek bir xlsx yazıp okur.
"""
import datetime

from openpyxl import Workbook, load_workbook

from excel_utils import SEMA_SAYFASI, excel_olustur, mevcut_verileri_oku


def _satir():
    return {"fatura_no": "GIB2024000000101", "sirket_adi": "ACME A.Ş.",
            "vkn": "1234567890", "vergiler_dahil_tutar": 1180.0,
            "kdv_haric_tutar": 1000.0, "para_birimi": "TRY",
            "fatura_tarihi": datetime.datetime(2024, 1, 15),
            "dosya_yolu": r"C:\x\a.pdf", "_teknik_bilgi": "Dijital"}


def _sutunlari_kaydir(yol, harita):
    """Faturalar sayfasının sütunlarını `harita` (eski → yeni) ile yeniden dizer.

    Kullanıcının ya da ileride şemanın sütun düzenini değiştirmesinin taklidi.
    """
    eski = load_workbook(yol)
    kaynak = eski["Faturalar"]
    yeni = Workbook()
    ws = yeni.active
    ws.title = "Faturalar"
    for row in kaynak.iter_rows():
        for c in row:
            if c.value is not None:
                ws.cell(row=c.row, column=harita.get(c.column, c.column), value=c.value)
    return eski, yeni


def test_yeni_dosyada_gizli_sema_sayfasi_var(tmp_path):
    yol = str(tmp_path / "f.xlsx")
    excel_olustur([_satir()], yol)

    wb = load_workbook(yol)
    ws = wb[SEMA_SAYFASI]
    assert ws.sheet_state == "hidden"
    harita = {r[0]: r[1] for r in ws.iter_rows(min_row=2, values_only=True)}
    # Bağımsız doğruluk kaynağı: bugünkü sabit düzen (bkz. tests/test_fatura.py).
    assert harita["fatura_no"] == 2
    assert harita["sirket_adi"] == 11
    assert harita["dosya_yolu_gizli"] == 15
    assert harita["kaynak"] == 16
    assert wb.active.title == "Faturalar"


def test_gomulu_sema_kaydirilmis_sutunlari_dogru_okur(tmp_path):
    """Sütunlar yer değiştirmiş, başlıklar da silinmiş: yalnızca gömülü şema kurtarır."""
    yol = str(tmp_path / "f.xlsx")
    excel_olustur([_satir()], yol)

    # fatura_no (2) ↔ sirket_adi (11), gizli yol (15) → 20
    kayma = {2: 11, 11: 2, 15: 20}
    eski, yeni = _sutunlari_kaydir(yol, kayma)
    for c in yeni["Faturalar"][1]:
        c.value = None                       # başlık eşlemesi devre dışı
    sema = yeni.create_sheet(SEMA_SAYFASI)
    sema.append(["anahtar", "sutun"])
    for anahtar, sutun in eski[SEMA_SAYFASI].iter_rows(min_row=2, values_only=True):
        sema.append([anahtar, kayma.get(sutun, sutun)])
    yeni.save(yol)

    satirlar, islenenmis = mevcut_verileri_oku(yol)
    assert satirlar[0]["fatura_no"] == "GIB2024000000101"
    assert satirlar[0]["sirket_adi"] == "ACME A.Ş."
    assert satirlar[0]["dosya_yolu"] == r"C:\x\a.pdf"
    assert islenenmis == {"a.pdf"}


def test_gomulu_semasiz_dosya_baslik_eslemesiyle_okunur(tmp_path):
    """Şema sayfası yok ama başlıklar duruyor: sütunlar kaysa da başlıktan bulunur."""
    yol = str(tmp_path / "f.xlsx")
    excel_olustur([_satir()], yol)

    # fatura_no (2) ↔ sirket_adi (11); başlıklar da birlikte taşınır
    _, yeni = _sutunlari_kaydir(yol, {2: 11, 11: 2})
    yeni.save(yol)

    satirlar, islenenmis = mevcut_verileri_oku(yol)
    assert satirlar[0]["fatura_no"] == "GIB2024000000101"
    assert satirlar[0]["sirket_adi"] == "ACME A.Ş."
    assert satirlar[0]["vkn"] == "1234567890"
    assert islenenmis == {"a.pdf"}


def test_ne_sema_ne_baslik_olan_eski_dosya_sabit_sutunlardan_okunur(tmp_path):
    """Başlık satırı tanınmıyor, şema sayfası yok: bugünkü sabit numaralar."""
    yol = str(tmp_path / "f.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Faturalar"
    ws.cell(row=1, column=1, value="eski başlık")
    ws.cell(row=2, column=2, value="GIB2024000000101")    # fatura_no
    ws.cell(row=2, column=11, value="ACME A.Ş.")          # sirket_adi
    ws.cell(row=2, column=7, value=1180.0)                # vergiler_dahil_tutar
    ws.cell(row=2, column=15, value=r"C:\x\a.pdf")        # gizli tam yol
    ws.cell(row=2, column=16, value="OCR")                # kaynak
    wb.save(yol)

    satirlar, islenenmis = mevcut_verileri_oku(yol)
    assert satirlar[0]["fatura_no"] == "GIB2024000000101"
    assert satirlar[0]["sirket_adi"] == "ACME A.Ş."
    assert satirlar[0]["vergiler_dahil_tutar"] == 1180.0
    assert satirlar[0]["_teknik_bilgi"] == "OCR"
    assert islenenmis == {"a.pdf"}

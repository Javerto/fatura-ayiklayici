"""Şema semantiği (#6): pasif alan, çekirdek kısıtı, zorunlu bayrağı.

Pasif alan = "kaldırılmış" alan. Kullanıcının tek tıkla yapabileceği bir işlem;
bu yüzden kritik test, pasif alanın arşiv verisinin bir sonraki onaydan sağ
çıkmasıdır.
"""
import datetime
import json
import queue
import threading

import pytest
from openpyxl import load_workbook

import sema
from ozet import ozet_hesapla
from api import Api
from excel_utils import excel_olustur, mevcut_verileri_oku
from extraction import veri_dogrula
from review import form_satira_uygula, satir_form_degerleri
from worker import worker
from tests.test_worker import SahteIstemci, _kuyrugu_bosalt, _pdf_olustur


def _sema(**degisiklik):
    """Tohumdan türetilmiş test şeması: anahtar → {alan: değer} değişiklikleri,
    `_ekle` ile yeni alanlar."""
    ham = sema.VARSAYILAN.json()
    for d in ham["alanlar"]:
        d.update(degisiklik.get(d["anahtar"], {}))
    ham["alanlar"] += degisiklik.get("_ekle", [])
    return sema.coz(ham)


def _eski_satir(no, sira_no=None):
    return {"fatura_no": no, "fatura_tarihi": datetime.datetime(2024, 1, 15),
            "sirket_adi": "ACME A.Ş.", "vkn": "1234567890",
            "vergiler_dahil_tutar": 1180.0, "kdv_haric_tutar": 1000.0,
            "para_birimi": "TRY", "sira_no": sira_no,
            "dosya_yolu": rf"C:\x\{no}.pdf", "_teknik_bilgi": "Dijital"}


def test_pasif_alanin_arsiv_verisi_sonraki_onaydan_sag_cikar(tmp_path):
    """Kullanıcı sira_no'yu "kaldırır" (pasife alır), yeni klasörü işler, onaylar.
    Eski faturaların sıra no'su Excel'de durmalı; yeni faturada doldurulmamalı."""
    kok, klasor = tmp_path / "ayar", tmp_path / "faturalar"
    kok.mkdir()
    klasor.mkdir()
    cikti = klasor / "cikti.xlsx"
    excel_olustur([_eski_satir("GIB2024000000001", sira_no=42.0)], str(cikti))

    Api(kok=kok)
    ham = json.loads((kok / "alanlar.json").read_text("utf-8"))
    next(d for d in ham["alanlar"] if d["anahtar"] == "sira_no")["aktif"] = False
    (kok / "alanlar.json").write_text(json.dumps(ham, ensure_ascii=False), "utf-8")
    api = Api(kok=kok)
    api._pencere = type("P", (), {"evaluate_js": lambda self, k: None})()

    _pdf_olustur(klasor / "yeni.pdf")
    # Model promptta istenmese de alanı döndürebilir — yine de doldurulmamalı.
    istemci = SahteIstemci(json.dumps({
        "fatura_no": "GIB2024000000002", "fatura_tarihi": "16.01.2024",
        "sirket_adi": "ACME A.Ş.", "vkn": "1234567890",
        "vergiler_dahil_tutar": 1180, "kdv_haric_tutar": 1000,
        "para_birimi": "TRY", "sira_no": 7}))
    log_q = queue.Queue()
    worker("K", str(klasor), "cikti.xlsx", log_q, threading.Event(),
           istemci=istemci, sema=api._sema)

    assert '"sira_no"' not in istemci.cagrilar[0][0]          # prompt'ta yok
    payload = next(d for t, d in _kuyrugu_bosalt(log_q) if t == "review")
    olay = api._review_olayi(payload)
    api._review = payload
    assert "sira_no" not in [a[0] for a in olay["alanlar"]]    # formda yok
    form = olay["satirlar"][0]["form"]
    assert "sira_no" not in form
    form["sirket_adi"] = "ACME DÜZELTİLMİŞ"                   # formdan bir düzenleme

    assert api.review_onayla({"0": form}, [], []).get("ok")
    satirlar, _ = mevcut_verileri_oku(str(cikti), sema=api._sema)
    sira = {s["fatura_no"]: s["sira_no"] for s in satirlar}
    assert sira == {"GIB2024000000001": 42.0, "GIB2024000000002": None}


def test_form_uygulamasi_pasif_alanin_degerine_dokunmaz():
    """Form pasif alanı taşımaz; uygulama onu "boş girildi" sayıp silmemeli."""
    s = _sema(sira_no={"aktif": False})
    satir = _eski_satir("GIB2024000000001", sira_no=42.0)
    form = satir_form_degerleri(satir, s)
    assert "sira_no" not in form
    assert form_satira_uygula(satir, form, s)["sira_no"] == 42.0


@pytest.mark.parametrize("anahtar", sema.CEKIRDEK)
def test_cekirdek_alan_pasife_alinamaz(anahtar):
    with pytest.raises(ValueError, match="çekirdek"):
        _sema(**{anahtar: {"aktif": False}})


def test_cekirdek_olmayan_tohum_alani_pasife_alinabilir():
    s = _sema(kdv_haric_tutar={"aktif": False})
    assert "kdv_haric_tutar" not in [a[0] for a in s.duzenlenebilir]
    assert "kdv_haric_tutar" in s.sutun                        # sütun korunur


def test_cekirdegi_pasif_dosya_tohuma_duser_ve_dosyaya_dokunulmaz(tmp_path):
    yol = tmp_path / "alanlar.json"
    ham = sema.VARSAYILAN.json()
    ham["alanlar"][0]["aktif"] = False                          # fatura_no
    yol.write_text(json.dumps(ham), "utf-8")
    assert sema.oku(yol) == sema.VARSAYILAN
    assert json.loads(yol.read_text("utf-8")) == ham


# ─── Doğrulama: zorunlu bayrağı, tip dönüşüm hatası, pasif alan ──────────────

_KULLANICI = {"_ekle": [
    {"etiket": "Sipariş No", "tip": "metin", "zorunlu": True},
    {"etiket": "Koli", "tip": "sayi"},
    {"etiket": "Vade", "tip": "tarih"},
]}


def _temiz(**ek):
    return {**_eski_satir("GIB2024000000001"), "siparis_no": "SP-1", **ek}


def test_zorunlu_kullanici_alani_bossa_uyarilir():
    s = _sema(**_KULLANICI)
    assert veri_dogrula(_temiz(), sema=s) == []
    assert [a for a, _ in veri_dogrula(_temiz(siparis_no=" "), sema=s)] == ["siparis_no"]


def test_zorunlu_olmayan_kullanici_alani_bossa_uyarilmaz():
    s = _sema(**_KULLANICI)
    assert veri_dogrula(_temiz(koli=None, vade=None), sema=s) == []


def test_formdaki_cevrilemeyen_sayi_ve_tarih_uyarilir():
    s = _sema(**_KULLANICI)
    form = {**satir_form_degerleri(_temiz(), s), "koli": "on iki", "vade": "yarın"}
    uyarilar = veri_dogrula(form_satira_uygula(_temiz(), form, s), sema=s)
    assert sorted(a for a, _ in uyarilar) == ["koli", "vade"]
    # Düzeltilince uyarı kalkar.
    form.update(koli="12", vade="01.02.2024")
    assert veri_dogrula(form_satira_uygula(_temiz(), form, s), sema=s) == []


def test_modelden_gelen_cevrilemeyen_sayi_uyarilir(tmp_path):
    s = _sema(**_KULLANICI)
    _pdf_olustur(tmp_path / "a.pdf")
    istemci = SahteIstemci(json.dumps({
        "fatura_no": "GIB2024000000002", "fatura_tarihi": "16.01.2024",
        "sirket_adi": "ACME", "vkn": "1234567890", "vergiler_dahil_tutar": 1180,
        "kdv_haric_tutar": 1000, "para_birimi": "TRY", "siparis_no": "SP-1",
        "koli": "bilinmiyor"}))
    log_q = queue.Queue()
    worker("K", str(tmp_path), "c.xlsx", log_q, threading.Event(),
           istemci=istemci, sema=s)
    olay = next(d for t, d in _kuyrugu_bosalt(log_q) if t == "fatura")
    assert [a for a, _ in olay["uyarilar"]] == ["koli"]


def test_pasif_alan_icin_uyari_uretilmez():
    """Pasif alan formda yok; uyarısı gösterilecek yer bulamazdı."""
    s = _sema(sira_no={"aktif": False},
              _ekle=[{"etiket": "Sipariş No", "tip": "metin",
                      "zorunlu": True, "aktif": False}])
    assert veri_dogrula(_temiz(sira_no=5000.0, siparis_no=None), sema=s) == []


def test_kdv_haric_pasifken_oran_kontrolu_atlanir():
    s = _sema(kdv_haric_tutar={"aktif": False})
    satir = _temiz(kdv_haric_tutar=1.0)            # aktif olsa saçma oran uyarısı
    assert veri_dogrula(satir) != []
    assert veri_dogrula(satir, sema=s) == []


def test_kdv_haric_pasifken_vergi_formulu_yeni_satirda_yazilmaz(tmp_path):
    """Yeni satırda hariç tutar yok: =dahil−hariç vergiyi dahil tutar gösterirdi.
    Eski satırın hariç tutarı arşivde duruyorsa formülü anlamlıdır, kalır."""
    s = _sema(kdv_haric_tutar={"aktif": False})
    yol = str(tmp_path / "f.xlsx")
    yeni = _temiz()
    del yeni["kdv_haric_tutar"]
    excel_olustur([_temiz(), yeni], yol, sema=s)
    ws = load_workbook(yol)["Faturalar"]
    vt = s.sutun["vergi_tutari"]
    assert str(ws.cell(row=2, column=vt).value).startswith("=")
    assert ws.cell(row=3, column=vt).value is None


def test_ozet_kullanici_alanlarindan_etkilenmez():
    s = [_temiz(), _temiz(koli=99.0, vade=datetime.datetime(2030, 1, 1))]
    assert ozet_hesapla(s) == ozet_hesapla([_temiz(), _temiz()])


def test_her_uyari_formdaki_bir_alana_denk_gelir():
    s = _sema(sira_no={"aktif": False}, **_KULLANICI)
    uyarilar = veri_dogrula({"sira_no": 5000.0, "dosya_yolu": "x.pdf"}, sema=s)
    assert uyarilar
    assert {a for a, _ in uyarilar} <= {a for a, _, _ in s.duzenlenebilir}

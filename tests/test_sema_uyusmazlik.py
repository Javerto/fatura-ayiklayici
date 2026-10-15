"""Şema/arşiv uyuşmazlık uyarısı (bilet #7).

Şema makine başına, arşiv klasör başınadır. Arşivin gömülü şeması aktif
şemadan farklıysa koşu düzeyinde uyarı çıkar; işlem durmaz.
"""
import datetime
import pathlib
import queue
import shutil
import threading

from openpyxl import load_workbook

import sema
from excel_utils import SEMA_SAYFASI, excel_olustur, sema_farki
from worker import worker


def _satir():
    return {"fatura_no": "GIB2024000000101", "sirket_adi": "ACME A.Ş.",
            "vkn": "1234567890", "vergiler_dahil_tutar": 1180.0,
            "kdv_haric_tutar": 1000.0, "para_birimi": "TRY",
            "fatura_tarihi": datetime.datetime(2024, 1, 15),
            "dosya_yolu": r"C:\x\a.pdf", "_teknik_bilgi": "Dijital"}


def _siparisli():
    ham = sema.VARSAYILAN.json()
    ham["alanlar"].append({"etiket": "Sipariş No", "tip": "metin",
                           "aciklama": "Sipariş numarası"})
    return sema.coz(ham)


def test_arsivde_olup_semada_olmayan_alan_sayilir(tmp_path):
    yol = str(tmp_path / "f.xlsx")
    excel_olustur([_satir()], yol, _siparisli())

    fark = sema_farki(yol, sema.VARSAYILAN)
    assert fark == {"arsivde_fazla": ["Sipariş No"], "yeni": []}


def test_semada_olup_arsivde_olmayan_alan_ayri_sayilir(tmp_path):
    yol = str(tmp_path / "f.xlsx")
    excel_olustur([_satir()], yol, sema.VARSAYILAN)

    fark = sema_farki(yol, _siparisli())
    assert fark == {"arsivde_fazla": [], "yeni": ["Sipariş No"]}


def test_semalar_ayniysa_uyari_yok(tmp_path):
    yol = str(tmp_path / "f.xlsx")
    excel_olustur([_satir()], yol, _siparisli())
    assert sema_farki(yol, _siparisli()) is None


def test_gomulu_semasi_olmayan_eski_arsivde_uyari_yok(tmp_path):
    yol = str(tmp_path / "f.xlsx")
    excel_olustur([_satir()], yol, _siparisli())
    wb = load_workbook(yol)
    del wb[SEMA_SAYFASI]                # bilet #4 öncesi sürümün yazdığı dosya
    wb.save(yol)
    assert sema_farki(yol, sema.VARSAYILAN) is None


def test_arsiv_yoksa_uyari_yok(tmp_path):
    assert sema_farki(str(tmp_path / "yok.xlsx"), _siparisli()) is None


def _review_yuku(klasor, aktif_sema):
    shutil.copy(pathlib.Path(__file__).parent / "fixtures" / "ornek_fatura.xml",
                klasor / "yeni.xml")
    log_q = queue.Queue()
    # XML-only akış modele gitmez; istemci yalnızca kurulmasın diye veriliyor.
    worker("ANAHTAR", str(klasor), "cikti.xlsx", log_q, threading.Event(),
           istemci=object(), sema=aktif_sema)
    olaylar = []
    while not log_q.empty():
        olaylar.append(log_q.get_nowait())
    return next(d for t, d in olaylar if t == "review")


def test_worker_uyusmazligi_review_yukune_koyar_ve_durmaz(tmp_path):
    excel_olustur([_satir()], str(tmp_path / "cikti.xlsx"), _siparisli())

    yuk = _review_yuku(tmp_path, sema.VARSAYILAN)
    assert yuk["sema_farki"] == {"arsivde_fazla": ["Sipariş No"], "yeni": []}
    assert len(yuk["yeni"]) == 1          # işlem durmadı, yeni fatura onaya geldi
    assert len(yuk["mevcut"]) == 1        # arşiv okunmaya devam etti


def test_uyusmazlik_onayi_engellemez(tmp_path):
    from api import Api

    cikti = str(tmp_path / "cikti.xlsx")
    excel_olustur([_satir()], cikti, _siparisli())
    api = Api(kok=tmp_path / "ayar")
    api._pencere = type("P", (), {"evaluate_js": lambda self, k: None})()
    api._review = _review_yuku(tmp_path, api._sema)
    assert api._review["sema_farki"]

    sonuc = api.review_onayla({}, [], [])
    assert sonuc["ok"] and sonuc["yazilan"] == 1
    # Arşive özgü sütun onaydan sağ çıkar, bu yüzden bilgi de sürer.
    assert sema_farki(cikti, api._sema) == {"arsivde_fazla": ["Sipariş No"], "yeni": []}


def test_arsive_ozgu_alanin_sutunu_ve_verisi_onaydan_sag_cikar(tmp_path):
    """Başka şemayla yazılmış arşiv aktif şemayla onaylanınca, aktif şemada
    olmayan sütun (başlığı, verisi, _Sema kaydı) sessizce siliniyordu."""
    from api import Api
    from excel_utils import mevcut_verileri_oku

    cikti = str(tmp_path / "cikti.xlsx")
    sip = _siparisli()
    excel_olustur([{**_satir(), "siparis_no": "SP-42"}], cikti, sip)
    api = Api(kok=tmp_path / "ayar")          # aktif şema: tohum, Sipariş No yok
    api._review = _review_yuku(tmp_path, api._sema)

    assert api.review_onayla({}, [], [])["ok"]

    wb = load_workbook(cikti)
    sutun = sip.sutun["siparis_no"]
    assert wb["Faturalar"].cell(1, sutun).value == "Sipariş No"
    assert wb["Faturalar"].cell(2, sutun).value == "SP-42"
    gomulu = dict(wb[SEMA_SAYFASI].iter_rows(min_row=2, max_col=2, values_only=True))
    assert gomulu["siparis_no"] == sutun
    # Bir sonraki onay da korur: değer arşivden yeniden okunuyor.
    assert mevcut_verileri_oku(cikti, sip)[0][0]["siparis_no"] == "SP-42"


def test_arsive_ozgu_alanin_sutunu_doluysa_kaybolmaz_tasinir(tmp_path):
    """Aktif şemadaki yeni bir alan arşive özgü sütunun numarasını almışsa
    arşiv verisi o sütunun altında ezilmez, boş bir sütuna taşınır."""
    from excel_utils import arsiv_semasi

    cikti = str(tmp_path / "cikti.xlsx")
    excel_olustur([{**_satir(), "siparis_no": "SP-42"}], cikti, _siparisli())
    ham = sema.VARSAYILAN.json()
    ham["alanlar"].append({"etiket": "Proje", "tip": "metin"})
    aktif = sema.coz(ham)                      # Proje de 17. sütunu alır
    assert aktif.sutun["proje"] == _siparisli().sutun["siparis_no"]

    s = arsiv_semasi(cikti, aktif)

    assert s.sutun["proje"] == 17 and s.sutun["siparis_no"] == 18
    excel_olustur([{**_satir(), "siparis_no": "SP-42", "proje": "P1"}],
                  cikti, s)
    ws = load_workbook(cikti)["Faturalar"]
    assert (ws.cell(1, 18).value, ws.cell(2, 18).value) == ("Sipariş No", "SP-42")
    assert ws.cell(2, 17).value == "P1"


def _vadeli_arsiv(cikti):
    ham = sema.VARSAYILAN.json()
    ham["alanlar"].append({"etiket": "Vade", "tip": "tarih"})
    vadeli = sema.coz(ham)
    excel_olustur([{**_satir(), "vade": datetime.datetime(2024, 2, 15)}], cikti, vadeli)
    return vadeli.sutun["vade"]


def _yeniden_yaz_tohumla(cikti):
    """Aktif şema tohum (Vade yok): arşiv okunup arşiv şemasıyla geri yazılır."""
    from excel_utils import arsiv_semasi, mevcut_verileri_oku
    s = arsiv_semasi(cikti, sema.VARSAYILAN)
    satirlar, _ = mevcut_verileri_oku(cikti, s)
    excel_olustur(satirlar, cikti, s)


def test_arsive_ozgu_tarih_sutunu_bicimini_korur(tmp_path):
    """arsiv_semasi tipi None ekliyordu: DD.MM.YYYY yerine yyyy-mm-dd h:mm:ss."""
    cikti = str(tmp_path / "cikti.xlsx")
    sutun = _vadeli_arsiv(cikti)

    _yeniden_yaz_tohumla(cikti)

    c = load_workbook(cikti)["Faturalar"].cell(2, sutun)
    assert c.value == datetime.datetime(2024, 2, 15)
    assert c.number_format == "DD.MM.YYYY"


def test_tipsiz_eski_sema_sayfasinda_tip_hucreden_cikarilir(tmp_path):
    cikti = str(tmp_path / "cikti.xlsx")
    sutun = _vadeli_arsiv(cikti)
    wb = load_workbook(cikti)
    wb[SEMA_SAYFASI].delete_cols(3, 10)      # eski biçim: yalnızca anahtar, sutun
    wb.save(cikti)

    _yeniden_yaz_tohumla(cikti)

    assert load_workbook(cikti)["Faturalar"].cell(2, sutun).number_format == "DD.MM.YYYY"


def test_worker_semalar_ayniysa_uyusmazlik_yok(tmp_path):
    excel_olustur([_satir()], str(tmp_path / "cikti.xlsx"), sema.VARSAYILAN)
    assert _review_yuku(tmp_path, sema.VARSAYILAN)["sema_farki"] is None

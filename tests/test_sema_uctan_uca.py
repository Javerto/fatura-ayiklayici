"""Şemaya elle eklenen alan uçtan uca: alanlar.json → prompt → model → form → Excel.

Kullanıcının yapacağının aynısı: dosyaya bir alan yazar, uygulamayı açar,
klasörü işler, gözden geçirmede düzeltir, onaylar. Ağ yok (sahte istemci),
gerçek alanlar.json yok (`Api(kok=tmp_path)`).
"""
import json
import queue
import threading

from openpyxl import load_workbook

from api import Api
from excel_utils import mevcut_verileri_oku
from worker import worker
from tests.test_worker import SahteIstemci, _kuyrugu_bosalt, _pdf_olustur


def _alan_ekle(kok, *alanlar):
    yol = kok / "alanlar.json"
    ham = json.loads(yol.read_text("utf-8"))
    ham["alanlar"] += list(alanlar)
    yol.write_text(json.dumps(ham, ensure_ascii=False), "utf-8")


def test_elle_eklenen_alan_promptan_excele_kadar_gider(tmp_path):
    kok, klasor = tmp_path / "ayar", tmp_path / "faturalar"
    kok.mkdir()
    klasor.mkdir()
    Api(kok=kok)                                  # ilk açılış: tohum yazılır
    assert (kok / "alanlar.json").exists()
    _alan_ekle(kok, {"etiket": "Sipariş No", "tip": "metin",
                     "aciklama": "alıcının sipariş numarası"})
    api = Api(kok=kok)                            # ikinci açılış: dosyadan okunur
    api._pencere = type("P", (), {"evaluate_js": lambda self, k: None})()

    _pdf_olustur(klasor / "a.pdf")
    istemci = SahteIstemci(json.dumps({
        "fatura_no": "GIB2024000000101", "fatura_tarihi": "15.01.2024",
        "sirket_adi": "ACME A.Ş.", "vkn": "1234567890",
        "vergiler_dahil_tutar": 1180, "kdv_haric_tutar": 1000,
        "para_birimi": "TRY", "siparis_no": "SP-77"}))
    log_q = queue.Queue()
    worker("K", str(klasor), "cikti.xlsx", log_q, threading.Event(),
           istemci=istemci, sema=api._sema)

    # prompt → model
    assert '"siparis_no": "alıcının sipariş numarası"' in istemci.cagrilar[0][0]
    payload = next(d for t, d in _kuyrugu_bosalt(log_q) if t == "review")
    assert payload["yeni"][0]["siparis_no"] == "SP-77"

    # form: alan gözden geçirme ekranında düzenlenebilir
    olay = api._review_olayi(payload)
    api._review = payload
    assert ["siparis_no", "Sipariş No", "metin"] in olay["alanlar"]
    form = olay["satirlar"][0]["form"]
    assert form["siparis_no"] == "SP-77"
    form["siparis_no"] = "SP-78"

    # Excel sütunu
    sonuc = api.review_onayla({"0": form}, [], [])
    assert sonuc.get("ok"), sonuc
    ws = load_workbook(payload["cikti"])["Faturalar"]
    assert ws.cell(row=1, column=17).value == "Sipariş No"
    assert ws.cell(row=2, column=17).value == "SP-78"
    satirlar, _ = mevcut_verileri_oku(payload["cikti"], sema=api._sema)
    assert satirlar[0]["siparis_no"] == "SP-78"


def test_ogrenilir_isaretli_eklenen_alan_firma_icin_hatirlanir(tmp_path):
    Api(kok=tmp_path)
    _alan_ekle(tmp_path, {"etiket": "Banka", "tip": "metin", "ogrenilir": True})
    api = Api(kok=tmp_path)
    api._review = {"mevcut": [], "atlanmis": [], "kesildi": False,
                   "cikti": str(tmp_path / "f.xlsx"),
                   "yeni": [{"fatura_no": "GIB2024000000101", "vkn": "1234567890",
                             "sirket_adi": "ACME", "banka": "Ziraat",
                             "dosya_yolu": str(tmp_path / "a.pdf")}]}
    api.review_onayla({}, [], [0])

    kurallar = json.loads((tmp_path / "duzeltmeler.json").read_text("utf-8"))
    assert kurallar["1234567890"]["banka"] == "Ziraat"
